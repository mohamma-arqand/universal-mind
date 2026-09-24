"""Tests: R44 item 10 — the formula-bearing Persian spreadsheet.

Live laws:
1. A Persian spreadsheet reads RIGHT-TO-LEFT (the view flag is on disk).
2. «جمع» earns a REAL =SUM formula — read back, the formula is ALIVE in the
   cell (never a fabricated computed number).
3. The production chain («... و در اکسل بریز و جمعش را بزن») ships the
   workbook with the formula row.
"""

from __future__ import annotations

from typing import Any


class TestTheFormulaSheet:
    """ExcelSuite.write_table — rtl + total."""

    def test_rtl_lands_on_disk(self) -> None:
        from openpyxl import load_workbook

        from universal_mind.excel_suite import ExcelSuite

        out = ExcelSuite().write_table(headers=["شاخص", "مقدار"], rows=[["a", 1]], rtl=True)
        ws = load_workbook(out["path"]).active
        assert ws.sheet_view.rightToLeft is True
        assert out["rtl"] is True

    def test_total_is_a_live_formula(self) -> None:
        from openpyxl import load_workbook

        from universal_mind.excel_suite import ExcelSuite

        out = ExcelSuite().write_table(
            headers=["شاخص", "مقدار"],
            rows=[["a", 10.0], ["b", 20.0], ["c", 30.0]],
            total=True,
        )
        assert out["formula_row"] == 5
        ws = load_workbook(out["path"]).active
        assert ws["B5"].value == "=SUM(B2:B4)"  # a REAL formula, not a number
        assert ws["A5"].value == "جمع"

    def test_no_total_no_formula_row(self) -> None:
        from universal_mind.excel_suite import ExcelSuite

        out = ExcelSuite().write_table(headers=["a", "b"], rows=[[1, 2]])
        assert out["formula_row"] is None

    def test_a_text_column_earns_no_sum(self) -> None:
        from openpyxl import load_workbook

        from universal_mind.excel_suite import ExcelSuite

        out = ExcelSuite().write_table(
            headers=["نام", "مقدار"], rows=[["الف", 1.0], ["ب", 2.0]], total=True,
        )
        ws = load_workbook(out["path"]).active
        assert ws["A4"].value == "جمع"  # label column untouched
        assert ws["B4"].value == "=SUM(B2:B3)"


class TestTheProductionChain:
    """The real route: numbers → excel with a live formula row."""

    def test_the_chain_ships_the_formula(self) -> None:
        import tempfile
        from pathlib import Path
        from unittest.mock import patch as mock_patch

        from openpyxl import load_workbook

        from universal_mind.database_suite import DatabaseSuite

        iso = DatabaseSuite(str(Path(tempfile.mkdtemp()) / "r44-10.db"))
        with mock_patch.object(DatabaseSuite, "shared_persistent",
                               classmethod(lambda cls: iso)):
            from universal_mind.persian_router import route_and_run

            p: Any = route_and_run("میانگین ۱۰ و ۲۰ و ۳۰ را حساب کن و در اکسل بریز و جمعش را بزن")
        assert p["route"] == ["data", "excel"]
        x = p["result"]["excel"]
        assert x["rtl"] is True and x["formula_row"] is not None

        ws = load_workbook(x["path"]).active
        assert ws.sheet_view.rightToLeft is True
        last = ws.max_row
        assert str(ws.cell(row=last, column=2).value).startswith("=SUM(")

    def test_read_back_shows_the_live_formula(self) -> None:
        import tempfile
        from pathlib import Path
        from unittest.mock import patch as mock_patch

        from universal_mind.database_suite import DatabaseSuite

        iso = DatabaseSuite(str(Path(tempfile.mkdtemp()) / "r44-10b.db"))
        with mock_patch.object(DatabaseSuite, "shared_persistent",
                               classmethod(lambda cls: iso)):
            from universal_mind.excel_suite import ExcelSuite
            from universal_mind.persian_router import route_and_run

            p = route_and_run("میانگین ۴ و ۶ را حساب کن و در اکسل بریز و جمعش را بزن")
            back = ExcelSuite().read_table(p["result"]["excel"]["path"])
        assert back["headers"] == ["شاخص", "مقدار"]
        assert any(
            isinstance(c, str) and c.startswith("=SUM(") for c in back["rows"][-1]
        )

    def test_without_jam_no_formula(self) -> None:
        import tempfile
        from pathlib import Path
        from unittest.mock import patch as mock_patch

        from universal_mind.database_suite import DatabaseSuite

        iso = DatabaseSuite(str(Path(tempfile.mkdtemp()) / "r44-10c.db"))
        with mock_patch.object(DatabaseSuite, "shared_persistent",
                               classmethod(lambda cls: iso)):
            from universal_mind.persian_router import route_and_run

            p = route_and_run("میانگین ۴ و ۶ را حساب کن و در اکسل بریز")
        assert p["result"]["excel"]["formula_row"] is None
