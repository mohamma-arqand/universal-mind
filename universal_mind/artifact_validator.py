"""The artifact validator — someone finally OPENS the files.

R46 item 1: a run that reports «سند PDF ساخته شد» with a byte count has
never proven the file is a REAL, OPENABLE artifact — until now. Every
successful run whose output carries a path gets a second, independent
check: the file is re-opened with the format's own reader and its shape
is asserted. A corrupted file fails with its NAME, and the run is
stamped `verified=0` — an honest red, never a silent maybe.

Supported proofs (the platform's own real outputs):
  .pdf   → starts with %PDF + pypdf opens it + >=1 page
  .png/.jpg → PIL.Image.open + verify()
  .xlsx  → openpyxl loads it + names its sheets
  .csv   → csv reader opens it + >=1 row
  .zip   → zipfile opens it + namelist() non-empty
Anything else: 'unverified' (named, not guessed).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


def validate_artifact(path_str: str) -> dict[str, Any]:
    """Open the file with its own format reader; return the proof."""
    p = Path(path_str)
    if not p.is_file():
        return {"ok": False, "kind": "missing", "detail": f"فایل نیست: {p.name}", "path": path_str}
    suffix = p.suffix.lower()
    try:
        if suffix == ".pdf":
            head = p.read_bytes()[:5]
            if head != b"%PDF-":
                return {"ok": False, "kind": "pdf", "detail": "سردر %PDF ندارد", "path": path_str}
            from pypdf import PdfReader

            n = len(PdfReader(str(p)).pages)
            if n < 1:
                return {"ok": False, "kind": "pdf", "detail": "صفحهای ندارد", "path": path_str}
            # R61-S6 — the page count is PERSIAN (a Latin digit in the
            # operator's report is a leak; review finding 13).
            _fa = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
            return {"ok": True, "kind": "pdf",
                    "detail": f"PDF سالم با {str(n).translate(_fa)} صفحه",
                    "path": path_str}
        if suffix in (".png", ".jpg", ".jpeg"):
            from PIL import Image

            with Image.open(p) as img:
                img.verify()
            return {"ok": True, "kind": "image", "detail": f"تصویر سالم {p.name}", "path": path_str}
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(str(p), read_only=True)
            names = list(wb.sheetnames)
            wb.close()
            if not names:
                return {"ok": False, "kind": "xlsx", "detail": "شیت ندارد", "path": path_str}
            return {"ok": True, "kind": "xlsx", "detail": f"اکسل سالم با شیت {names[0]}", "path": path_str}
        if suffix == ".csv":
            import csv

            with p.open(newline="", encoding="utf-8-sig") as fh:
                rows = list(csv.reader(fh))
            if not rows:
                return {"ok": False, "kind": "csv", "detail": "ردیف ندارد", "path": path_str}
            width = len(rows[0])
            ragged = sum(1 for r in rows if len(r) != width)
            if ragged:
                return {"ok": False, "kind": "csv",
                        "detail": f"CSV کج است — {ragged} ردیف با عرض نادرست", "path": path_str}
            return {"ok": True, "kind": "csv", "detail": f"CSV سالم با {len(rows)} ردیف", "path": path_str}
        if suffix == ".zip":
            import zipfile

            with zipfile.ZipFile(p) as zf:
                names = zf.namelist()
            if not names:
                return {"ok": False, "kind": "zip", "detail": "خالی است", "path": path_str}
            bad = zipfile.ZipFile(str(p)).testzip()
            if bad:
                return {"ok": False, "kind": "zip", "detail": f"عضو خراب: {bad}", "path": path_str}
            return {"ok": True, "kind": "zip", "detail": f"ZIP سالم با {len(names)} عضو", "path": path_str}
    except Exception as exc:  # noqa: BLE001 — the reader's own verdict is the proof
        return {"ok": False, "kind": suffix.lstrip(".") or "?", "detail": f"باز نشد: {exc}", "path": path_str}
    return {"ok": False, "kind": "unverified", "detail": f"قالبِ شناختهشده نیست ({suffix})", "path": path_str}


def _find_artifacts(payload: dict[str, Any]) -> list[str]:
    """The real output paths of a run payload (recursive, path-shaped)."""
    found: list[str] = []

    def _walk(node: Any) -> None:
        if isinstance(node, dict):
            for k, v in node.items():
                if k in ("path", "pdf", "image_path", "out", "out_path", "file", "zip_path", "xlsx_path") and isinstance(v, str):
                    found.append(v)
                else:
                    _walk(v)
        elif isinstance(node, (list, tuple)):
            for item in node:
                _walk(item)

    _walk(payload)
    return [f for f in found if f]


def validate_run(payload: dict[str, Any]) -> dict[str, Any]:
    """Validate EVERY artifact of a run; the stamp rides into the report."""
    paths = _find_artifacts(payload.get("result") or {})
    if not paths:
        return {"ok": True, "checked": 0, "verdicts": [], "report": "", "error": ""}
    verdicts = [validate_artifact(p) for p in paths]
    good = [v for v in verdicts if v["ok"]]
    bad = [v for v in verdicts if not v["ok"] and v["kind"] != "unverified"]
    parts = [v["detail"] for v in good]
    if bad:
        parts += [f"⚠️ {Path(v['path']).name}: {v['detail']}" for v in bad]
    report = " | ".join(parts)
    return {
        "ok": not bad,
        "checked": len(verdicts),
        "verdicts": verdicts,
        "report": f"تأییدِ فایل: {report}" if report else "",
        "error": "",
    }
