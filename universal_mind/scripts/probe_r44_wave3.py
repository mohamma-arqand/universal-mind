#!/usr/bin/env python3
"""Probe: R44 wave-3 — the ears, the formula sheet, the message.

Live laws:
1. «گوش کن» is a real GESTURE: the heard text is echoed and routed (the
   echo precedes the run — a misheard command never executes silently),
   and a missing engine / silence is honest.
2. «... و جمعش را بزن» ships a Persian RTL workbook with a REAL live
   =SUM formula in the cell (read back from disk, not assumed).
3. «... و به آدرس x@y ایمیل کن» produces a REAL .eml with the chain's
   artifact attached; SMTP without credentials says so instead of faking.
"""

from __future__ import annotations

import sys

sys.path.insert(0, "..")

sys.stderr.write("PROBE R44 wave-3 (ears + formula sheet + message):\n")


def _ok(name: str, cond: bool, extra: str = "") -> None:
    mark = "PASS" if cond else "FAIL"
    sys.stderr.write(f"  [{mark}] {name}" + (f" — {extra}" if extra else "") + "\n")
    if not cond:
        raise SystemExit(1)


def main() -> int:
    import tempfile
    from pathlib import Path
    from unittest.mock import patch as mock_patch

    from universal_mind.database_suite import DatabaseSuite

    outbox = Path(tempfile.mkdtemp()) / "outbox"
    outbox.mkdir(parents=True, exist_ok=True)
    iso = DatabaseSuite(str(Path(tempfile.mkdtemp()) / "probe-w3.db"))

    with mock_patch.object(DatabaseSuite, "shared_persistent", classmethod(lambda cls: iso)), \
         mock_patch("universal_mind.email_outbox._outbox_dir", lambda *_a, **_k: outbox):
        from universal_mind.persian_router import route_and_run

        # H1 — the ears: heard text is echoed, then really routed.
        with mock_patch(
            "universal_mind.speech_tool.SpeechTool.listen",
            return_value={"ok": True, "recognized": "میانگین ۴ و ۶ را حساب کن",
                          "lang": "fa-IR", "seconds": 5, "note": "", "error": ""},
        ):
            heard = route_and_run("گوش کن")
        _ok("the heard command really ran", heard.get("route") == ["data"], f"route={heard.get('route')}")
        _ok("the echo precedes the run", "شنیدم" in str(heard.get("agent_report")))

        with mock_patch(
            "universal_mind.speech_tool.SpeechTool.listen",
            return_value={"ok": False, "recognized": "", "error": "میکروفون در دسترس نیست"},
        ):
            deaf = route_and_run("گوش کن")
        _ok("a deaf platform never invents work",
            deaf.get("ok") is False and "نشنیدم" in str(deaf.get("agent_report")))

        # H2 — the formula sheet: RTL on disk + a live =SUM in the cell.
        from openpyxl import load_workbook

        sheet = route_and_run("میانگین ۱۰ و ۲۰ و ۳۰ را حساب کن و در اکسل بریز و جمعش را بزن")
        x = sheet["result"]["excel"]
        ws = load_workbook(x["path"]).active
        _ok("the sheet reads right-to-left", ws.sheet_view.rightToLeft is True)
        last = ws.max_row
        formula = str(ws.cell(row=last, column=2).value)
        _ok("a REAL formula lives in the cell", formula.startswith("=SUM("), formula)

        # H3 — the message: a real .eml with the chain's artifact attached.
        mail = route_and_run(
            "میانگین ۱۰ و ۲۰ را حساب کن و گزارشش کن و به آدرس ali@example.com ایمیل کن"
        )
        em = mail["result"]["email"]
        eml_files = list(outbox.glob("*.eml"))
        _ok("a real .eml landed on disk", bool(eml_files), f"{len(eml_files)} file(s)")
        _ok("the chain's artifact really attached", int(em.get("attachment_bytes") or 0) > 1000,
            f"{em.get('attachment_bytes')} bytes")
        raw = eml_files[0].read_bytes()
        _ok("the message is RFC-822 shaped",
            b"To: ali@example.com" in raw and b"MIME-Version: 1.0" in raw)

        # SMTP is optional and honest (no credentials in the environment here).
        from universal_mind.email_outbox import send

        no_smtp = send(str(eml_files[0].path if hasattr(eml_files[0], "path") else eml_files[0]))
        _ok("no credentials -> an honest 'not sent'", no_smtp.get("sent") is False
            and "UM_SMTP_HOST" in str(no_smtp.get("error")))

    sys.stderr.write("R44 wave-3: ALL HOLDS GREEN\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
