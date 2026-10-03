"""R74 P2 — the FORMAT CONVERTER (the 30th real capability).

«فایل X را به CSV/اکسل/JSON تبدیل کن» had no owner: CSV fell to a data
mean, Excel crashed openpyxl on a .txt (not a named refusal), JSON was
unknown. The converter reads the REAL source (csv/txt/json/xlsx), builds
the REAL target with the stdlib/openpyxl, and names both sides. A source
the target format cannot represent is REFUSED BY NAME — never a crash.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

_READERS = (".txt", ".csv", ".json", ".xlsx", ".md")
_WRITERS = (".csv", ".xlsx", ".json", ".txt")


def _read_rows(src: Path) -> dict[str, Any]:
    """Read the source into headers+rows; every unreadable fact is named."""
    suffix = src.suffix.lower()
    if suffix in (".txt", ".md"):
        text = src.read_text(encoding="utf-8", errors="replace")
        lines = [ln for ln in text.splitlines() if ln.strip()]
        if not lines:
            return {"ok": False, "error": f"فایل خالی است: {src}"}
        # delimiter sniff: comma/semicolon/tab, else each line one cell
        delim = ","
        if lines[0].count(";") > lines[0].count(","):
            delim = ";"
        elif "\t" in lines[0] and lines[0].count("\t") > lines[0].count(","):
            delim = "\t"
        rows = [ln.split(delim) for ln in lines]
        return {"ok": True, "headers": rows[0], "rows": rows[1:] or [], "count": len(rows)}
    if suffix == ".csv":
        with src.open(encoding="utf-8-sig", newline="") as fh:
            got = list(csv.reader(fh))
        if not got:
            return {"ok": False, "error": f"فایل CSV خالی است: {src}"}
        return {"ok": True, "headers": got[0], "rows": got[1:], "count": len(got)}
    if suffix == ".json":
        try:
            data = json.loads(src.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            return {"ok": False, "error": f"JSON ناسالم: {e}"}
        if isinstance(data, list) and data and isinstance(data[0], dict):
            return {"ok": True, "headers": list(data[0].keys()),
                    "rows": [[r.get(h, "") for h in data[0].keys()] for r in data],
                    "count": len(data)}
        return {"ok": False, "error": "JSON یک آرایهٔ اشیاء نیست — قالب پشتیبانیشده نیست"}
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            return {"ok": False, "error": "openpyxl روی این ماشین نصب نیست"}
        wb = load_workbook(src, data_only=False)
        ws = wb.active
        grid = [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows()]
        if not grid:
            return {"ok": False, "error": f"برگه خالی است: {src}"}
        return {"ok": True, "headers": grid[0], "rows": grid[1:], "count": len(grid)}
    return {"ok": False, "error": f"منبعِ «{suffix}» خوانده نمیشود — قالبهای خواندنی: {', '.join(_READERS)}"}


def convert_format(src: str, target_fmt: str) -> dict[str, Any]:
    """Convert a REAL file to csv/xlsx/json — both sides named, never invented."""
    s = Path(src)
    if not s.exists():
        return {"ok": False, "error": f"فایل پیدا نشد: {src}"}
    fmt = target_fmt.lower().lstrip(".")
    if fmt not in ("csv", "xlsx", "excel", "json"):
        return {"ok": False,
                "error": f"قالب مقصد «{target_fmt}» پشتیبانی نمیشود — مقصدها: csv، اکسل (xlsx)، json"}
    if fmt == "excel":
        fmt = "xlsx"
    if s.suffix.lower() == f".{fmt}":
        return {"ok": False, "error": f"منبع از قبل {fmt.upper()} است — تبدیل به خودش بیمعناست"}

    got = _read_rows(s)
    if not got.get("ok"):
        return {"ok": False, "error": str(got.get("error"))}
    headers, rows = got["headers"], got["rows"]

    out = s.with_suffix(f".{fmt}")
    if fmt == "csv":
        with out.open("w", encoding="utf-8-sig", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(headers)
            w.writerows(rows)
    elif fmt == "json":
        payload = [dict(zip([str(h) for h in headers], [str(c) for c in r])) for r in rows]
        out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    else:  # xlsx
        try:
            from openpyxl import Workbook
        except ImportError:
            return {"ok": False, "error": "openpyxl روی این ماشین نصب نیست"}
        wb = Workbook()
        ws = wb.active
        ws.title = "داده"
        ws.append([str(h) for h in headers])
        for r in rows:
            ws.append([str(c) for c in r])
        wb.save(out)

    return {"ok": True, "src": str(s), "path": str(out),
            "rows": len(rows), "bytes": out.stat().st_size, "error": ""}
